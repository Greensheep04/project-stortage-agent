#!/usr/bin/env python3
"""T-004b 完整集程序化校验。

在试样校验基础上扩展：试样原样并入核对、明细覆盖与单明细上限、
场景组数、近似候选清单、manifest 与分批记录。
"""

import hashlib
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from t004.text_utils import has_intra_duplicate, sentences

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
PILOT = PROJECT_ROOT / "deliverables" / "t004" / "pilot-v2"
FULL = PROJECT_ROOT / "deliverables" / "t004" / "full-v2"
XLSX = PROJECT_ROOT / "deliverables" / "仓储异常测试数据.xlsx"
XLSX_SHA256 = "61edff98759f2697c5489d2dd2fd4962637ea38982601d219e6b40a37ddd1df1"
AS_OF = datetime.fromisoformat("2026-09-09T20:00:00+08:00")

FORBIDDEN_LINES = [
    ("SO-1009", "001"), ("SO-1009", "002"), ("SO-1039", "002"), ("SO-1042", "001"),
    ("SO-1050", "001"), ("SO-1050", "002"), ("SO-1075", "001"), ("SO-1077", "001"),
    ("SO-1083", "001"), ("SO-1099", "001"),
]
SCENARIO_MIN = {
    "insufficient_evidence": 10,
    "same_object_time_contradiction": 10,
    "compatible_state_change": 10,
    "similar_history_interference": 10,
    "sop_version_boundary": 10,
    "backfill_time_visibility": 10,
    "same_order_different_line": 10,
}
EXPECTED_COUNTS = {"event": 700, "case": 280, "sop": 20}
MIN_COVERED_LINES = 100
MAX_EVENTS_PER_LINE = 10

COMMON_FIELDS = ["doc_id", "doc_type", "title", "recorded_at", "author_role", "scope", "sections", "synthetic"]
FILE_FOR_TYPE = {"event": "events.jsonl", "case": "cases.jsonl", "sop": "sops.jsonl"}

FAILURES = []
WARNINGS = []
PASSES = []


def ok(msg):
    PASSES.append(msg)


def check(cond, msg):
    if not cond:
        FAILURES.append(msg)
    return cond


def warn(cond, msg):
    if not cond:
        WARNINGS.append(msg)
    return cond


def parse_ts(value, label):
    try:
        return datetime.fromisoformat(value)
    except (TypeError, ValueError):
        FAILURES.append(f"{label}: 时间格式无法解析: {value!r}")
        return None


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def load_jsonl(path):
    items = []
    for lineno, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        try:
            items.append(json.loads(line))
        except json.JSONDecodeError as e:
            FAILURES.append(f"{path.name}:{lineno} JSON 解析失败: {e}")
    return items


def raw_lines(path):
    return [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def load_source_rows():
    import openpyxl

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    col_names = {
        "发货计划": ("订单号", "订单明细号"),
        "仓储流水": ("订单号", "订单明细号"),
        "业务说明": ("关联订单号", "关联订单明细号"),
    }
    rows = {}
    for sheet, (order_col, line_col) in col_names.items():
        ws = wb[sheet]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        oi, li = header.index(order_col), header.index(line_col)
        sheet_rows = {}
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            sheet_rows[r] = (row[oi], row[li])
        rows[sheet] = sheet_rows
    return rows


def check_common(doc, file_name):
    doc_id = doc.get("doc_id", "?")
    label = f"{file_name}:{doc_id}"
    for field in COMMON_FIELDS:
        check(field in doc, f"{label} 缺少字段 {field}")
    check(isinstance(doc.get("doc_id"), str) and doc.get("doc_id"), f"{label} doc_id 非字符串")
    check(doc.get("doc_type") in FILE_FOR_TYPE, f"{label} doc_type 非法: {doc.get('doc_type')!r}")
    check(FILE_FOR_TYPE.get(doc.get("doc_type")) == file_name, f"{label} doc_type 与所在文件不一致")
    check(isinstance(doc.get("title"), str) and doc.get("title"), f"{label} title 为空")
    check(doc.get("synthetic") is True, f"{label} synthetic 必须为 true")
    scope = doc.get("scope")
    if isinstance(scope, dict):
        for key in ("warehouse_id", "sku", "order_id", "line_id"):
            check(key in scope, f"{label} scope 缺少 {key}")
    else:
        FAILURES.append(f"{label} scope 非对象")
    sections = doc.get("sections")
    if isinstance(sections, list) and sections:
        ids = []
        for section in sections:
            if not isinstance(section, dict):
                FAILURES.append(f"{label} section 非对象")
                continue
            for key in ("section_id", "heading", "text"):
                check(key in section, f"{label} section 缺少 {key}")
            if isinstance(section.get("text"), str):
                check(bool(section["text"].strip()), f"{label} section 文本为空")
            ids.append(section.get("section_id"))
        check(len(ids) == len(set(ids)), f"{label} section_id 重复")
    else:
        FAILURES.append(f"{label} sections 为空或非数组")


def check_event(doc, file_name, source_rows):
    label = f"{file_name}:{doc.get('doc_id')}"
    occurred = parse_ts(doc.get("occurred_at"), f"{label} occurred_at")
    recorded = parse_ts(doc.get("recorded_at"), f"{label} recorded_at")
    if occurred and recorded:
        check(occurred <= recorded, f"{label} occurred_at 晚于 recorded_at")
        check(recorded <= AS_OF, f"{label} recorded_at 晚于数据截止时间")
    check(isinstance(doc.get("transaction_ids"), list), f"{label} transaction_ids 非数组")
    refs = doc.get("source_refs")
    if not isinstance(refs, list) or not refs:
        FAILURES.append(f"{label} source_refs 为空")
        return
    scope = doc.get("scope") or {}
    for ref in refs:
        sheet, row = ref.get("sheet"), ref.get("row")
        check(ref.get("file") == "deliverables/仓储异常测试数据.xlsx", f"{label} source_refs file 非预期")
        if sheet not in source_rows:
            FAILURES.append(f"{label} source_refs sheet 非法: {sheet!r}")
            continue
        if row not in source_rows[sheet]:
            FAILURES.append(f"{label} source_refs 行不存在: {sheet} 第 {row} 行")
            continue
        ref_order, ref_line = source_rows[sheet][row]
        check(
            str(ref_order) == str(scope.get("order_id")) and str(ref_line) == str(scope.get("line_id")),
            f"{label} source_refs {sheet} 第 {row} 行与 scope 不一致",
        )


def check_case(doc, file_name, sop_by_id):
    label = f"{file_name}:{doc.get('doc_id')}"
    closed = parse_ts(doc.get("closed_at"), f"{label} closed_at")
    recorded = parse_ts(doc.get("recorded_at"), f"{label} recorded_at")
    if closed and recorded:
        check(closed <= recorded, f"{label} closed_at 晚于 recorded_at")
        check(recorded <= AS_OF, f"{label} recorded_at 晚于数据截止时间")
    scope = doc.get("scope") or {}
    check(scope.get("order_id") is None, f"{label} 案例 scope.order_id 应为 null")
    check(scope.get("line_id") is None, f"{label} 案例 scope.line_id 应为 null")
    check(isinstance(doc.get("history_case_id"), str) and doc["history_case_id"], f"{label} history_case_id 缺失")
    refs = doc.get("rule_refs")
    if not isinstance(refs, list):
        FAILURES.append(f"{label} rule_refs 非数组")
        return
    for ref in refs:
        sop = sop_by_id.get(ref)
        if sop is None:
            FAILURES.append(f"{label} rule_refs 指向不存在的 SOP: {ref}")
            continue
        if closed:
            start = parse_ts(sop["effective_from"], f"{ref} effective_from")
            end = parse_ts(sop["effective_to"], f"{ref} effective_to") if sop.get("effective_to") else None
            check(start <= closed and (end is None or closed < end), f"{label} rule_refs {ref} 在结案时点不适用")


def check_sop(doc, file_name):
    label = f"{file_name}:{doc.get('doc_id')}"
    recorded = parse_ts(doc.get("recorded_at"), f"{label} recorded_at")
    start = parse_ts(doc.get("effective_from"), f"{label} effective_from")
    end = parse_ts(doc.get("effective_to"), f"{label} effective_to") if doc.get("effective_to") else None
    if recorded and start:
        check(recorded <= start, f"{label} recorded_at 晚于 effective_from")
    if start and end:
        check(start < end, f"{label} effective_from 不早于 effective_to")
    scope = doc.get("scope") or {}
    check(scope.get("order_id") is None, f"{label} SOP scope.order_id 应为 null")
    check(scope.get("line_id") is None, f"{label} SOP scope.line_id 应为 null")
    check(isinstance(doc.get("policy_id"), str) and doc["policy_id"], f"{label} policy_id 缺失")
    check(isinstance(doc.get("version"), str) and doc["version"], f"{label} version 缺失")


MONTH_WORDS = {
    "一月": 1, "二月": 2, "三月": 3, "四月": 4, "五月": 5, "六月": 6,
    "七月": 7, "八月": 8, "九月": 9, "十月": 10, "十一月": 11, "十二月": 12,
}


def load_source_times():
    import openpyxl

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    times = {}
    id_times = {}
    id_rows = {}
    for sheet, tcol, icol in (
        ("发货计划", None, None),
        ("仓储流水", "发生时间", "流水号"),
        ("业务说明", "记录时间", "证据编号"),
    ):
        ws = wb[sheet]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        ti = header.index(tcol) if tcol else None
        ii = header.index(icol) if icol else None
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            t = row[ti] if ti is not None else None
            times[(sheet, r)] = t
            if ii is not None and row[ii]:
                id_times[str(row[ii])] = t
                id_rows[str(row[ii])] = (sheet, r)
    return times, id_times, id_rows


def check_event_time_visibility(events, src_times, id_times):
    import re as _re

    ref_violations, text_violations = [], []
    for e in events:
        rec = parse_ts(e.get("recorded_at"), e.get("doc_id", "?"))
        if rec is not None and rec.tzinfo is not None:
            rec = rec.replace(tzinfo=None)
        for ref in e.get("source_refs", []):
            t = src_times.get((ref.get("sheet"), ref.get("row")))
            if t and rec and t > rec:
                ref_violations.append(f"{e.get('doc_id')}:{ref.get('sheet')}#{ref.get('row')}")
        text = "".join(s.get("text", "") for s in e.get("sections", []))
        for mid in set(_re.findall(r"(?:TX|EV)-\d+", text)):
            t = id_times.get(mid)
            if t and rec and t > rec:
                text_violations.append(f"{e.get('doc_id')}:{mid}")
    return ref_violations, text_violations


def check_case_months(cases):
    violations = []
    for c in cases:
        text = "".join(s.get("text", "") for s in c.get("sections", []))
        months = [m for w, m in MONTH_WORDS.items() if w in text]
        if not months:
            continue
        closed = parse_ts(c.get("closed_at"), c.get("doc_id", "?"))
        if closed and max(months) > closed.month:
            violations.append(f"{c.get('doc_id')}:月份{max(months)}>结案{closed.month}")
    return violations


def self_test():
    src_times = {("仓储流水", 2): datetime(2026, 9, 8, 12, 0)}
    id_times = {"TX-1": datetime(2026, 9, 8, 12, 0)}
    bad_event = {
        "doc_id": "EVX-TEST",
        "recorded_at": "2026-09-08T10:00:00+08:00",
        "source_refs": [{"sheet": "仓储流水", "row": 2}],
        "sections": [{"text": "TX-1 已过账。"}],
    }
    ref_v, text_v = check_event_time_visibility([bad_event], src_times, id_times)
    assert ref_v and text_v, (ref_v, text_v)
    bad_case = {
        "doc_id": "CASE-TEST",
        "closed_at": "2026-06-15T10:00:00+08:00",
        "sections": [{"text": "七月的一次装运。"}],
    }
    assert check_case_months([bad_case]), "月份倒置未检出"
    print("self-test: 通过（未来引用与月份倒置均被检出）")


def bigrams(text):
    s = "".join(ch for ch in text if not ch.isspace())
    return {s[i:i + 2] for i in range(len(s) - 1)}


def main():
    if "--self-test" in sys.argv:
        self_test()
        return 0
    events = load_jsonl(FULL / "events.jsonl")
    cases = load_jsonl(FULL / "cases.jsonl")
    sops = load_jsonl(FULL / "sops.jsonl")
    scenarios = load_jsonl(FULL / "review" / "scenarios.jsonl")

    check(sha256_file(XLSX) == XLSX_SHA256, "源 Excel 哈希与合同基线不一致")

    for doc in events:
        check_common(doc, "events.jsonl")
    for doc in cases:
        check_common(doc, "cases.jsonl")
    for doc in sops:
        check_common(doc, "sops.jsonl")

    all_docs = events + cases + sops
    doc_ids = [d.get("doc_id") for d in all_docs]
    check(len(doc_ids) == len(set(doc_ids)), "doc_id 存在重复")
    check(all(not str(i).startswith("EXAMPLE") for i in doc_ids), "交付中出现 EXAMPLE 前缀")
    for doc_type, expected in EXPECTED_COUNTS.items():
        actual = sum(1 for d in all_docs if d.get("doc_type") == doc_type)
        check(actual == expected, f"{doc_type} 篇数 {actual} != {expected}")

    # 试样原样并入
    for name, pilot_file, full_file in (
        ("events", PILOT / "events.jsonl", FULL / "events.jsonl"),
        ("cases", PILOT / "cases.jsonl", FULL / "cases.jsonl"),
        ("sops", PILOT / "sops.jsonl", FULL / "sops.jsonl"),
        ("scenarios", PILOT / "review" / "scenarios.jsonl", FULL / "review" / "scenarios.jsonl"),
    ):
        pilot_raw = raw_lines(pilot_file)
        full_raw = raw_lines(full_file)
        check(full_raw[:len(pilot_raw)] == pilot_raw, f"试样 {name} 未原样并入")

    source_rows = load_source_rows()
    for doc in events:
        check_event(doc, "events.jsonl", source_rows)

    sop_by_id = {d["doc_id"]: d for d in sops}
    for doc in cases:
        check_case(doc, "cases.jsonl", sop_by_id)
    for doc in sops:
        check_sop(doc, "sops.jsonl")

    # SOP 区间与版本差异
    groups = defaultdict(list)
    for doc in sops:
        scope = doc["scope"]
        groups[(doc["policy_id"], scope.get("warehouse_id"), scope.get("sku"))].append(doc)
    for key, docs in groups.items():
        ordered = sorted(docs, key=lambda d: parse_ts(d["effective_from"], d["doc_id"]) or AS_OF)
        for prev, nxt in zip(ordered, ordered[1:]):
            prev_to = parse_ts(prev["effective_to"], prev["doc_id"]) if prev.get("effective_to") else None
            nxt_from = parse_ts(nxt["effective_from"], nxt["doc_id"])
            check(prev_to is not None and prev_to <= nxt_from, f"政策 {key} 区间重叠: {prev['doc_id']}/{nxt['doc_id']}")
    by_policy = defaultdict(list)
    for doc in sops:
        by_policy[doc["policy_id"]].append(doc)
    for policy, docs in by_policy.items():
        check(len(docs) == 2, f"政策 {policy} 版本数 {len(docs)} != 2")
        if len(docs) < 2:
            continue
        a, b = docs
        text_a = "".join(s["text"].strip() for s in a["sections"])
        text_b = "".join(s["text"].strip() for s in b["sections"])
        diff_sections = sum(
            1 for sa in a["sections"]
            if not any(sb["section_id"] == sa["section_id"] and sb["text"] == sa["text"] for sb in b["sections"])
        )
        check(text_a != text_b, f"{policy} 两版全文相同")
        check(diff_sections >= 2, f"{policy} 两版实质差异不足（仅 {diff_sections} 节不同）")
        ok(f"政策 {policy} 两版实质差异通过")
    for doc in sops:
        sup = doc.get("supersedes")
        if sup is None:
            continue
        target = sop_by_id.get(sup)
        if target is None:
            FAILURES.append(f"{doc['doc_id']} supersedes 指向不存在文档 {sup}")
            continue
        check(target["policy_id"] == doc["policy_id"], f"{doc['doc_id']} supersedes 政策不一致")
        check(target.get("effective_to") == doc["effective_from"], f"{doc['doc_id']} supersedes 区间未衔接")

    # 禁补明细结构
    forbidden_set = {f"{o}/{l}" for o, l in FORBIDDEN_LINES}
    forbidden_reports = []
    for doc in events:
        scope = doc.get("scope") or {}
        key = f"{scope.get('order_id')}/{scope.get('line_id')}"
        if key not in forbidden_set:
            continue
        section_ids = {s.get("section_id") for s in doc["sections"]}
        headings = " ".join(str(s.get("heading", "")) for s in doc["sections"])
        has_pending = bool(section_ids & {"pending", "待查", "open"}) or any(w in headings for w in ("待查", "待办", "未确认"))
        has_cause = bool(section_ids & {"cause", "result", "conclusion"})
        check(has_pending, f"禁补明细 {key} 的 {doc['doc_id']} 缺少待查节")
        check(not has_cause, f"禁补明细 {key} 的 {doc['doc_id']} 出现原因/结论节")
        forbidden_reports.append(f"{key}:{doc['doc_id']}")
    covered_forbidden = {r.split(":")[0] for r in forbidden_reports}
    check(covered_forbidden == forbidden_set, f"10 个禁补明细未全部覆盖: 缺 {sorted(forbidden_set - covered_forbidden)}")
    ok(f"10 个禁补明细均有待查结构，共 {len(forbidden_reports)} 篇事件")

    # 明细覆盖与单明细上限
    per_line = Counter(f"{e['scope']['order_id']}/{e['scope']['line_id']}" for e in events)
    check(len(per_line) >= MIN_COVERED_LINES, f"事件覆盖明细 {len(per_line)} < {MIN_COVERED_LINES}")
    over = {k: v for k, v in per_line.items() if v > MAX_EVENTS_PER_LINE}
    check(not over, f"单明细事件超过 {MAX_EVENTS_PER_LINE} 篇: {over}")
    ok(f"事件覆盖明细 {len(per_line)} 个，单明细上限通过（最大 {max(per_line.values())} 篇）")

    # 场景检查
    scenario_ids = [s.get("scenario_id") for s in scenarios]
    check(len(scenario_ids) == len(set(scenario_ids)), "scenario_id 重复")
    type_counts = Counter()
    required = ["scenario_id", "scenario_type", "related_doc_ids", "order_id", "line_id", "cutoff",
                "facts_to_check", "allowed_conclusions", "forbidden_conclusions", "basis"]
    for scenario in scenarios:
        sid = scenario.get("scenario_id", "?")
        for field in required:
            check(field in scenario, f"场景 {sid} 缺少字段 {field}")
        type_counts[scenario.get("scenario_type")] += 1
        for ref in scenario.get("related_doc_ids", []):
            check(ref in doc_ids, f"场景 {sid} 引用不存在的文档 {ref}")
        parse_ts(scenario.get("cutoff"), f"场景 {sid} cutoff")
    for scenario_type, minimum in SCENARIO_MIN.items():
        check(type_counts.get(scenario_type, 0) >= minimum,
              f"场景类型 {scenario_type} 数量 {type_counts.get(scenario_type, 0)} < {minimum}")
    ok(f"场景总数 {len(scenarios)}，类型分布 {dict(type_counts)}")

    # T-004c 时间一致性检查
    src_times, id_times, _ = load_source_times()
    ref_violations, text_violations = check_event_time_visibility(events, src_times, id_times)
    check(not ref_violations, f"事件引用未来来源: {ref_violations[:5]}（共 {len(ref_violations)}）")
    check(not text_violations, f"事件正文提及未来事实: {text_violations[:5]}（共 {len(text_violations)}）")
    if not ref_violations and not text_violations:
        ok("事件时间可见性：未来引用 = 0，正文未来事实 = 0")
    month_violations = check_case_months(cases)
    check(not month_violations, f"案例月份倒置: {month_violations[:5]}")
    if not month_violations:
        ok("案例月份/时间倒置 = 0")

    # 文本质量检查（R1/R2）
    dup_docs = [
        d["doc_id"]
        for d in all_docs
        if any(
            has_intra_duplicate(s["text"], sim_threshold=2.0 if d["doc_type"] == "sop" else 0.5)
            for s in d["sections"]
        )
    ]
    check(not dup_docs, f"句内重复未清零: {dup_docs[:5]}（共 {len(dup_docs)} 篇）")
    if not dup_docs:
        ok("句内重复（含包含关系）= 0")

    case_sent = Counter()
    for c in cases:
        seen = set()
        for s in c["sections"]:
            seen.update(sentences(s["text"]))
        for x in seen:
            case_sent[x] += 1
    over40 = [(s, n) for s, n in case_sent.items() if n > 40]
    check(not over40, f"案例整句超过 40 篇共用: {over40[:3]}")
    ok(f"案例整句最高共用 {case_sent.most_common(1)[0][1] if case_sent else 0} 篇（上限 40）")

    # 同主题变体：标题唯一 + 两两相似度上限
    by_theme = defaultdict(list)
    for c in cases:
        theme = c["title"].split("案例（")[0]
        by_theme[theme].append(c)
    theme_max_sim = 0.0
    sim_violations = []
    for theme, docs in by_theme.items():
        titles = [d["title"] for d in docs]
        check(len(titles) == len(set(titles)), f"主题 {theme} 标题不唯一")
        grams = [bigrams("".join(s["text"] for s in d["sections"])) for d in docs]
        for i in range(len(grams)):
            for j in range(i + 1, len(grams)):
                a, b = grams[i], grams[j]
                if not a or not b:
                    continue
                sim = len(a & b) / len(a | b)
                theme_max_sim = max(theme_max_sim, sim)
                if sim >= 0.8:
                    sim_violations.append((theme, docs[i]["doc_id"], docs[j]["doc_id"], round(sim, 3)))
    check(not sim_violations, f"同主题变体相似度过高（≥0.8）: {sim_violations[:3]}")
    ok(f"同主题变体最大两两相似度 {theme_max_sim:.3f}（阈值 <0.8），主题数 {len(by_theme)}")

    # 近似候选清单
    def near_duplicates(docs, label, threshold=0.85, cap=50):
        grams = [(d["doc_id"], bigrams("".join(s["text"] for s in d["sections"]))) for d in docs]
        pairs = []
        for i in range(len(grams)):
            for j in range(i + 1, len(grams)):
                a, b = grams[i][1], grams[j][1]
                if not a or not b:
                    continue
                jac = len(a & b) / len(a | b)
                if jac >= threshold:
                    pairs.append((round(jac, 3), grams[i][0], grams[j][0]))
        pairs.sort(reverse=True)
        ok(f"{label} 近似候选（≥{threshold}）{len(pairs)} 对，列出前 {min(cap, len(pairs))} 对供中层判读")
        return pairs[:cap]

    event_pairs = near_duplicates(events, "事件")
    case_pairs = near_duplicates(cases, "案例")
    (FULL / "review").mkdir(parents=True, exist_ok=True)
    (FULL / "review" / "near-duplicates.json").write_text(
        json.dumps({"events": event_pairs, "cases": case_pairs}, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # manifest
    manifest_path = FULL / "manifest.json"
    if check(manifest_path.exists(), "manifest.json 缺失"):
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        files = {f["name"]: f for f in manifest.get("files", [])}
        for name, count in (
            ("events.jsonl", len(events)), ("cases.jsonl", len(cases)),
            ("sops.jsonl", len(sops)), ("review/scenarios.jsonl", len(scenarios)),
        ):
            entry = files.get(name)
            if entry is None:
                FAILURES.append(f"manifest 缺少 {name}")
                continue
            check(entry.get("count") == count, f"manifest {name} 篇数不一致")
            check(entry.get("sha256") == sha256_file(FULL / name), f"manifest {name} sha256 不一致")
        check(manifest.get("source_snapshot", {}).get("sha256") == XLSX_SHA256, "manifest 源快照哈希不一致")
        check(manifest.get("generation", {}).get("dashscope_api_calls") == 0, "manifest DashScope 调用不为 0")
        check(bool(manifest.get("generation", {}).get("batches")), "manifest 缺少分批记录")
        ok(f"manifest 分批记录 {len(manifest.get('generation', {}).get('batches', []))} 批")

    print("== T-004b 完整集程序化校验 ==")
    print(f"PASS {len(PASSES)} 项")
    for item in PASSES:
        print(f"  [PASS] {item}")
    print(f"WARN {len(WARNINGS)} 项")
    for item in WARNINGS:
        print(f"  [WARN] {item}")
    print(f"FAIL {len(FAILURES)} 项")
    for item in FAILURES:
        print(f"  [FAIL] {item}")
    print("禁补明细事件:", len(forbidden_reports))
    return 1 if FAILURES else 0


if __name__ == "__main__":
    raise SystemExit(main())
