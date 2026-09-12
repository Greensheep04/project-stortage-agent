#!/usr/bin/env python3
"""T-004a 试样程序化校验。

检查 JSONL 格式、必填与类型、doc_id 唯一、source_refs 与只读工作簿一致、
时间约束、SOP 版本区间与实质差异、10 个禁补明细结构、场景组数量与引用、
manifest 一致性。输出 PASS/WARN/FAIL 报告，失败时退出码非零。
"""

import hashlib
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
PILOT_DIR = PROJECT_ROOT / "deliverables" / "t004" / "pilot-v1"
XLSX = PROJECT_ROOT / "deliverables" / "仓储异常测试数据.xlsx"
XLSX_SHA256 = "61edff98759f2697c5489d2dd2fd4962637ea38982601d219e6b40a37ddd1df1"
AS_OF = datetime.fromisoformat("2026-09-09T20:00:00+08:00")

FORBIDDEN_LINES = [
    ("SO-1009", "001"),
    ("SO-1009", "002"),
    ("SO-1039", "002"),
    ("SO-1042", "001"),
    ("SO-1050", "001"),
    ("SO-1050", "002"),
    ("SO-1075", "001"),
    ("SO-1077", "001"),
    ("SO-1083", "001"),
    ("SO-1099", "001"),
]

SCENARIO_MIN = {
    "insufficient_evidence": 3,
    "same_object_time_contradiction": 2,
    "compatible_state_change": 2,
    "similar_history_interference": 2,
    "sop_version_boundary": 2,
    "backfill_time_visibility": 2,
    "same_order_different_line": 2,
}

COMMON_FIELDS = [
    "doc_id",
    "doc_type",
    "title",
    "recorded_at",
    "author_role",
    "scope",
    "sections",
    "synthetic",
]
EXTRA_FIELDS = {
    "event": ["occurred_at", "transaction_ids", "source_refs"],
    "case": ["history_case_id", "closed_at", "rule_refs"],
    "sop": ["policy_id", "version", "effective_from", "effective_to", "supersedes"],
}
FILE_FOR_TYPE = {"event": "events.jsonl", "case": "cases.jsonl", "sop": "sops.jsonl"}
EXPECTED_COUNTS = {"event": 24, "case": 12, "sop": 4}
PENDING_SECTION_IDS = {"pending", "待查", "open"}

FAILURES = []
WARNINGS = []
PASSES = []


def ok(msg):
    PASSES.append(msg)


def check(cond, msg):
    if not cond:
        FAILURES.append(msg)
    return cond


def warn(cond, msg):
    if not cond:
        WARNINGS.append(msg)
    return cond


def parse_ts(value, label):
    try:
        return datetime.fromisoformat(value)
    except (TypeError, ValueError):
        FAILURES.append(f"{label}: 时间格式无法解析: {value!r}")
        return None


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def load_jsonl(path):
    items = []
    for lineno, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        try:
            items.append(json.loads(line))
        except json.JSONDecodeError as e:
            FAILURES.append(f"{path.name}:{lineno} JSON 解析失败: {e}")
    return items


def load_source_rows():
    import openpyxl

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    col_names = {
        "发货计划": ("订单号", "订单明细号"),
        "仓储流水": ("订单号", "订单明细号"),
        "业务说明": ("关联订单号", "关联订单明细号"),
    }
    rows = {}
    for sheet, (order_col, line_col) in col_names.items():
        ws = wb[sheet]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        oi, li = header.index(order_col), header.index(line_col)
        sheet_rows = {}
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            sheet_rows[r] = (row[oi], row[li])
        rows[sheet] = sheet_rows
    return rows


def check_common(doc, file_name):
    doc_id = doc.get("doc_id", "?")
    label = f"{file_name}:{doc_id}"
    for field in COMMON_FIELDS:
        check(field in doc, f"{label} 缺少字段 {field}")
    check(isinstance(doc.get("doc_id"), str) and doc.get("doc_id"), f"{label} doc_id 非字符串")
    check(
        doc.get("doc_type") in FILE_FOR_TYPE,
        f"{label} doc_type 非法: {doc.get('doc_type')!r}",
    )
    check(
        FILE_FOR_TYPE.get(doc.get("doc_type")) == file_name,
        f"{label} doc_type 与所在文件不一致",
    )
    check(isinstance(doc.get("title"), str) and doc.get("title"), f"{label} title 为空")
    check(doc.get("synthetic") is True, f"{label} synthetic 必须为 true")
    scope = doc.get("scope")
    if isinstance(scope, dict):
        for key in ("warehouse_id", "sku", "order_id", "line_id"):
            check(key in scope, f"{label} scope 缺少 {key}")
    else:
        FAILURES.append(f"{label} scope 非对象")
    sections = doc.get("sections")
    if isinstance(sections, list) and sections:
        ids = []
        for section in sections:
            if not isinstance(section, dict):
                FAILURES.append(f"{label} section 非对象")
                continue
            for key in ("section_id", "heading", "text"):
                check(key in section, f"{label} section 缺少 {key}")
            if isinstance(section.get("text"), str):
                check(bool(section["text"].strip()), f"{label} section 文本为空")
            ids.append(section.get("section_id"))
        check(len(ids) == len(set(ids)), f"{label} section_id 重复")
    else:
        FAILURES.append(f"{label} sections 为空或非数组")


def check_event(doc, file_name, source_rows):
    label = f"{file_name}:{doc.get('doc_id')}"
    occurred = parse_ts(doc.get("occurred_at"), f"{label} occurred_at")
    recorded = parse_ts(doc.get("recorded_at"), f"{label} recorded_at")
    if occurred and recorded:
        check(occurred <= recorded, f"{label} occurred_at 晚于 recorded_at")
        check(recorded <= AS_OF, f"{label} recorded_at 晚于数据截止时间")
    txs = doc.get("transaction_ids")
    check(isinstance(txs, list), f"{label} transaction_ids 非数组")
    refs = doc.get("source_refs")
    if not isinstance(refs, list) or not refs:
        FAILURES.append(f"{label} source_refs 为空")
        return
    scope = doc.get("scope") or {}
    for ref in refs:
        sheet = ref.get("sheet")
        row = ref.get("row")
        check(
            ref.get("file") == "deliverables/仓储异常测试数据.xlsx",
            f"{label} source_refs file 非预期: {ref.get('file')!r}",
        )
        if sheet not in source_rows:
            FAILURES.append(f"{label} source_refs sheet 非法: {sheet!r}")
            continue
        if row not in source_rows[sheet]:
            FAILURES.append(f"{label} source_refs 行不存在: {sheet} 第 {row} 行")
            continue
        ref_order, ref_line = source_rows[sheet][row]
        check(
            str(ref_order) == str(scope.get("order_id"))
            and str(ref_line) == str(scope.get("line_id")),
            f"{label} source_refs {sheet} 第 {row} 行的订单/明细与 scope 不一致",
        )


def check_case(doc, file_name, sop_by_id):
    label = f"{file_name}:{doc.get('doc_id')}"
    closed = parse_ts(doc.get("closed_at"), f"{label} closed_at")
    recorded = parse_ts(doc.get("recorded_at"), f"{label} recorded_at")
    if closed and recorded:
        check(closed <= recorded, f"{label} closed_at 晚于 recorded_at")
        check(recorded <= AS_OF, f"{label} recorded_at 晚于数据截止时间")
    scope = doc.get("scope") or {}
    check(scope.get("order_id") is None, f"{label} 案例 scope.order_id 应为 null")
    check(scope.get("line_id") is None, f"{label} 案例 scope.line_id 应为 null")
    check(
        isinstance(doc.get("history_case_id"), str) and doc["history_case_id"],
        f"{label} history_case_id 缺失",
    )
    refs = doc.get("rule_refs")
    if not isinstance(refs, list):
        FAILURES.append(f"{label} rule_refs 非数组")
        return
    for ref in refs:
        sop = sop_by_id.get(ref)
        if sop is None:
            FAILURES.append(f"{label} rule_refs 指向不存在的 SOP: {ref}")
            continue
        if closed:
            from_ts = parse_ts(sop["effective_from"], f"{ref} effective_from")
            to_ts = parse_ts(sop["effective_to"], f"{ref} effective_to") if sop.get("effective_to") else None
            check(
                from_ts is not None and from_ts <= closed and (to_ts is None or closed < to_ts),
                f"{label} rule_refs {ref} 在结案时点不适用",
            )


def check_sop(doc, file_name):
    label = f"{file_name}:{doc.get('doc_id')}"
    recorded = parse_ts(doc.get("recorded_at"), f"{label} recorded_at")
    from_ts = parse_ts(doc.get("effective_from"), f"{label} effective_from")
    to_ts = parse_ts(doc.get("effective_to"), f"{label} effective_to") if doc.get("effective_to") else None
    if recorded and from_ts:
        check(recorded <= from_ts, f"{label} recorded_at 晚于 effective_from")
    if from_ts and to_ts:
        check(from_ts < to_ts, f"{label} effective_from 不早于 effective_to")
    scope = doc.get("scope") or {}
    check(scope.get("order_id") is None, f"{label} SOP scope.order_id 应为 null")
    check(scope.get("line_id") is None, f"{label} SOP scope.line_id 应为 null")
    check(isinstance(doc.get("policy_id"), str) and doc["policy_id"], f"{label} policy_id 缺失")
    check(isinstance(doc.get("version"), str) and doc["version"], f"{label} version 缺失")


def check_sop_intervals(sops):
    groups = defaultdict(list)
    for doc in sops:
        scope = doc["scope"]
        key = (doc["policy_id"], scope.get("warehouse_id"), scope.get("sku"))
        groups[key].append(doc)
    for key, docs in groups.items():
        ordered = sorted(docs, key=lambda d: parse_ts(d["effective_from"], d["doc_id"]) or AS_OF)
        for prev, nxt in zip(ordered, ordered[1:]):
            prev_to = parse_ts(prev["effective_to"], prev["doc_id"]) if prev.get("effective_to") else None
            nxt_from = parse_ts(nxt["effective_from"], nxt["doc_id"])
            if prev_to is None or (nxt_from and prev_to > nxt_from):
                FAILURES.append(f"政策 {key} 的区间重叠: {prev['doc_id']} 与 {nxt['doc_id']}")
            else:
                ok(f"政策 {key} 区间不重叠: {prev['doc_id']} → {nxt['doc_id']}")


def check_sop_supersedes(sops, sop_by_id):
    for doc in sops:
        sup = doc.get("supersedes")
        if sup is None:
            continue
        target = sop_by_id.get(sup)
        if target is None:
            FAILURES.append(f"{doc['doc_id']} supersedes 指向不存在文档 {sup}")
            continue
        check(target["policy_id"] == doc["policy_id"], f"{doc['doc_id']} supersedes 政策不一致")
        check(target["version"] != doc["version"], f"{doc['doc_id']} supersedes 版本相同")
        if target.get("effective_to"):
            check(
                target["effective_to"] == doc["effective_from"],
                f"{doc['doc_id']} supersedes 区间未衔接",
            )


def check_sop_substantive_diff(sops):
    by_policy = defaultdict(list)
    for doc in sops:
        by_policy[doc["policy_id"]].append(doc)
    for policy, docs in by_policy.items():
        if len(docs) < 2:
            WARNINGS.append(f"政策 {policy} 只有一个版本，无法比较差异")
            continue
        for i in range(len(docs)):
            for j in range(i + 1, len(docs)):
                a, b = docs[i], docs[j]
                text_a = "".join(s["text"].strip() for s in a["sections"])
                text_b = "".join(s["text"].strip() for s in b["sections"])
                diff_sections = sum(
                    1
                    for sa in a["sections"]
                    if not any(
                        sb["section_id"] == sa["section_id"] and sb["text"] == sa["text"]
                        for sb in b["sections"]
                    )
                )
                check(text_a != text_b, f"{policy} 版本 {a['doc_id']}/{b['doc_id']} 全文相同")
                check(
                    diff_sections >= 2,
                    f"{policy} 版本 {a['doc_id']}/{b['doc_id']} 实质差异不足（仅 {diff_sections} 节不同）",
                )
                if text_a != text_b and diff_sections >= 2:
                    ok(f"{policy} 版本 {a['doc_id']}/{b['doc_id']} 存在实质差异")


def main():
    events = load_jsonl(PILOT_DIR / "events.jsonl")
    cases = load_jsonl(PILOT_DIR / "cases.jsonl")
    sops = load_jsonl(PILOT_DIR / "sops.jsonl")
    scenarios = load_jsonl(PILOT_DIR / "review" / "scenarios.jsonl")

    check(
        sha256_file(XLSX) == XLSX_SHA256,
        f"源 Excel 哈希与合同基线不一致（期望 {XLSX_SHA256[:12]}…）",
    )

    for doc in events:
        check_common(doc, "events.jsonl")
    for doc in cases:
        check_common(doc, "cases.jsonl")
    for doc in sops:
        check_common(doc, "sops.jsonl")

    all_docs = events + cases + sops
    doc_ids = [d.get("doc_id") for d in all_docs]
    check(len(doc_ids) == len(set(doc_ids)), "doc_id 存在重复")
    check(
        all(not str(i).startswith("EXAMPLE") for i in doc_ids),
        "交付中出现 EXAMPLE 前缀",
    )

    for doc_type, expected in EXPECTED_COUNTS.items():
        actual = sum(1 for d in all_docs if d.get("doc_type") == doc_type)
        check(actual == expected, f"{doc_type} 篇数 {actual} != {expected}")

    source_rows = load_source_rows()
    for doc in events:
        check_event(doc, "events.jsonl", source_rows)

    sop_by_id = {d["doc_id"]: d for d in sops}
    for doc in cases:
        check_case(doc, "cases.jsonl", sop_by_id)
    for doc in sops:
        check_sop(doc, "sops.jsonl")
    check_sop_intervals(sops)
    check_sop_supersedes(sops, sop_by_id)
    check_sop_substantive_diff(sops)

    # 10 个禁补明细：只允许待查结构，且逐条列出供人工复核
    forbidden_reports = []
    forbidden_set = {f"{o}/{l}" for o, l in FORBIDDEN_LINES}
    for doc in events:
        scope = doc.get("scope") or {}
        key = f"{scope.get('order_id')}/{scope.get('line_id')}"
        if key not in forbidden_set:
            continue
        section_ids = {s.get("section_id") for s in doc["sections"]}
        headings = " ".join(str(s.get("heading", "")) for s in doc["sections"])
        has_pending = bool(section_ids & PENDING_SECTION_IDS) or any(
            word in headings for word in ("待查", "待办", "未确认")
        )
        has_cause = bool(section_ids & {"cause", "result", "conclusion"})
        check(has_pending, f"禁补明细 {key} 的 {doc['doc_id']} 缺少待查节")
        check(not has_cause, f"禁补明细 {key} 的 {doc['doc_id']} 出现原因/结论节")
        forbidden_reports.append(
            {
                "line": key,
                "doc_id": doc["doc_id"],
                "sections": [s["section_id"] for s in doc["sections"]],
            }
        )
    ok(f"10 个禁补明细结构检查完成，命中事件 {len(forbidden_reports)} 条（清单见下）")

    # 场景检查
    scenario_ids = [s.get("scenario_id") for s in scenarios]
    check(len(scenario_ids) == len(set(scenario_ids)), "scenario_id 重复")
    type_counts = defaultdict(int)
    scenario_required = [
        "scenario_id",
        "scenario_type",
        "related_doc_ids",
        "order_id",
        "line_id",
        "cutoff",
        "facts_to_check",
        "allowed_conclusions",
        "forbidden_conclusions",
        "basis",
    ]
    for scenario in scenarios:
        sid = scenario.get("scenario_id", "?")
        for field in scenario_required:
            check(field in scenario, f"场景 {sid} 缺少字段 {field}")
        type_counts[scenario.get("scenario_type")] += 1
        for ref in scenario.get("related_doc_ids", []):
            check(ref in doc_ids, f"场景 {sid} 引用不存在的文档 {ref}")
        parse_ts(scenario.get("cutoff"), f"场景 {sid} cutoff")
    for scenario_type, minimum in SCENARIO_MIN.items():
        check(
            type_counts.get(scenario_type, 0) >= minimum,
            f"场景类型 {scenario_type} 数量 {type_counts.get(scenario_type, 0)} < {minimum}",
        )
    ok(f"场景总数 {len(scenarios)}，类型分布 {dict(type_counts)}")

    # 禁补明细在证据不足场景中的覆盖
    insufficient_lines = {
        f"{s.get('order_id')}/{s.get('line_id')}"
        for s in scenarios
        if s.get("scenario_type") == "insufficient_evidence"
    }
    warn(
        len(insufficient_lines & forbidden_set) >= 3,
        f"证据不足场景覆盖禁补明细 {len(insufficient_lines & forbidden_set)} 条（建议至少 3）",
    )

    # manifest 一致性
    manifest_path = PILOT_DIR / "manifest.json"
    if check(manifest_path.exists(), "manifest.json 缺失"):
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        files = {f["name"]: f for f in manifest.get("files", [])}
        for name, count in (
            ("events.jsonl", len(events)),
            ("cases.jsonl", len(cases)),
            ("sops.jsonl", len(sops)),
            ("review/scenarios.jsonl", len(scenarios)),
        ):
            entry = files.get(name)
            if entry is None:
                FAILURES.append(f"manifest 缺少 {name}")
                continue
            check(entry.get("count") == count, f"manifest {name} 篇数不一致")
            check(
                entry.get("sha256") == sha256_file(PILOT_DIR / name),
                f"manifest {name} sha256 不一致",
            )
        check(
            manifest.get("source_snapshot", {}).get("sha256") == XLSX_SHA256,
            "manifest 源快照哈希不一致",
        )
        check(
            manifest.get("generation", {}).get("dashscope_api_calls") == 0,
            "manifest 记录的 DashScope API 调用不为 0",
        )

    print("== T-004a 试样程序化校验 ==")
    print(f"PASS {len(PASSES)} 项")
    for item in PASSES:
        print(f"  [PASS] {item}")
    print(f"WARN {len(WARNINGS)} 项")
    for item in WARNINGS:
        print(f"  [WARN] {item}")
    print(f"FAIL {len(FAILURES)} 项")
    for item in FAILURES:
        print(f"  [FAIL] {item}")
    if forbidden_reports:
        print("禁补明细事件清单（供人工复核）:")
        for item in forbidden_reports:
            print(f"  - {item['line']}: {item['doc_id']} sections={item['sections']}")

    return 1 if FAILURES else 0


if __name__ == "__main__":
    raise SystemExit(main())
