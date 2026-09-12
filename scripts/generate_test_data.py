import json
import random
import datetime
import collections
import sys
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_TEMPLATE = BASE_DIR / "inputs" / "仓储异常测试模板.xlsx"
OUTPUT_XLSX = BASE_DIR / "deliverables" / "仓储异常测试数据.xlsx"
OUTPUT_JSON = BASE_DIR / "deliverables" / "评测答案.json"

DEADLINE = datetime.datetime(2026, 9, 9, 20, 0, 0)
ORDER_START = datetime.datetime(2026, 9, 1, 8, 0, 0)
ORDER_END = datetime.datetime(2026, 9, 9, 18, 0, 0)
NOT_DUE_START = DEADLINE + datetime.timedelta(hours=2)
NOT_DUE_END = datetime.datetime(2026, 9, 11, 18, 0, 0)

SKUS = [
    ("MUG-BL", "陶瓷杯蓝色"),
    ("MUG-WH", "陶瓷杯白色"),
    ("BAG-BK", "帆布袋黑色"),
    ("BAG-KH", "帆布袋卡其色"),
]

REQUIRED_DISTRIBUTION = {
    "matched": 90,
    "split_matched": 20,
    "quantity_difference": 10,
    "insufficient_evidence": 10,
    "sku_mismatch": 5,
    "draft_only": 5,
    "reversal_matched": 5,
    "not_due": 5,
}

PLAN_HEADERS = [
    "订单号", "订单明细号", "计划仓库", "SKU", "商品名称", "应发数量", "单位",
    "下单时间", "约定出库时间", "订单状态", "来源系统", "业务备注"
]
FLOW_HEADERS = [
    "流水号", "仓储单号", "仓储明细号", "订单号", "订单明细号", "仓库", "SKU",
    "商品名称", "业务类型", "数量", "单位", "发生时间", "记录状态", "关联流水号",
    "来源系统", "仓储备注"
]
EXPLANATION_HEADERS = [
    "证据编号", "关联订单号", "关联订单明细号", "关联流水号", "记录时间",
    "记录类型", "记录人", "内容"
]

MUST_NOT_CLAIM = {
    "matched": ["存在数量差异", "存在SKU不一致", "分批出库"],
    "split_matched": ["重复出库", "数量差异", "SKU不一致"],
    "quantity_difference": ["已经补发", "客户已签收全部", "已经报废"],
    "insufficient_evidence": ["已经报废", "客户拒收", "仓库漏发"],
    "sku_mismatch": ["该SKU已正常发出", "差异仅由缺货导致", "数量无差异"],
    "draft_only": ["货物实际未发", "已正常出库", "已过账"],
    "reversal_matched": ["原多录流水已删除", "实际多发", "存在数量差异"],
    "not_due": ["已逾期", "仓库违约", "客户已签收"],
}

MATCHED_NOTES = [
    "按拣货单完成交接。",
    "扫描复核后装车发出。",
    "交接清点无误，已过账。",
    "出库实物与单据核对一致。",
    "拣货复核完成，条码已扫描。",
    "装车交接完成，单据已签收。",
]

FIRST_BATCH_NOTES = [
    "第一车装运{a}件。",
    "首批装车{a}件。",
    "上午交接{a}件。",
]

SECOND_BATCH_NOTES = [
    "第二车装运{b}件。",
    "剩余{b}件下午装车。",
    "下午交接{b}件。",
]

INSUFFICIENT_NOTES = [
    "按拣货单完成交接。",
    "扫描复核后过账。",
    "交接清点完成，单据已交接。",
    "装车发出，条码已扫描。",
    "出库复核完成。",
    "按单拣货并扫描过账。",
]


def rand_dt(start, end):
    if start >= end:
        return end
    delta_seconds = int((end - start).total_seconds())
    seconds = random.randrange(0, delta_seconds + 1, 60)
    return start + datetime.timedelta(seconds=seconds)


def next_flow_id(seq):
    seq[0] += 1
    return f"TX-{seq[0]:04d}"


def next_evidence_id(seq):
    seq[0] += 1
    return f"EV-{seq[0]:04d}"


def build_lines():
    lines = []
    for i in range(100):
        order_id = f"SO-{1001 + i}"
        n_lines = 2 if i < 50 else 1
        order_time = rand_dt(ORDER_START, ORDER_END)
        for l in range(n_lines):
            line_id = f"{l + 1:03d}"
            lines.append({
                "order_id": order_id,
                "line_id": line_id,
                "order_time": order_time,
            })

    categories = []
    for cat, count in REQUIRED_DISTRIBUTION.items():
        categories.extend([cat] * count)
    random.shuffle(categories)

    for line, cat in zip(lines, categories):
        line["category"] = cat
        sku_code, sku_name = random.choice(SKUS)
        line["sku"] = sku_code
        line["sku_name"] = sku_name

        if cat == "split_matched":
            planned = random.randint(6, 16)
        elif cat in ("quantity_difference", "insufficient_evidence"):
            planned = random.randint(5, 16)
        elif cat == "reversal_matched":
            planned = random.randint(3, 12)
        else:
            planned = random.randint(4, 15)
        line["planned"] = planned

        if cat == "not_due":
            agreed_time = rand_dt(NOT_DUE_START, NOT_DUE_END)
        else:
            earliest_agreed = line["order_time"] + datetime.timedelta(minutes=30)
            if earliest_agreed > DEADLINE:
                earliest_agreed = line["order_time"]
            agreed_time = rand_dt(earliest_agreed, DEADLINE)
        line["agreed_time"] = agreed_time
        line["order_status"] = "生效"
        line["source"] = "测试OMS"
        line["plan_note"] = ""

    return lines


def generate_flows_and_answers(lines):
    flow_seq = [1000]
    evidence_seq = [1000]
    flows = []
    explanations = []
    cases = []

    for line in lines:
        cat = line["category"]
        planned = line["planned"]
        order_id = line["order_id"]
        line_id = line["line_id"]
        sku = line["sku"]
        sku_name = line["sku_name"]
        order_time = line["order_time"]
        agreed_time = line["agreed_time"]

        def pick_time(start, end):
            return rand_dt(start, end)

        line_flows = []
        line_explanations = []

        if cat == "matched":
            t = pick_time(order_time, agreed_time)
            fid = next_flow_id(flow_seq)
            line_flows.append({
                "flow_id": fid,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "销售出库",
                "quantity": planned,
                "unit": "件",
                "time": t,
                "status": "已过账",
                "related_flow": None,
                "source": "测试WMS",
                "note": random.choice(MATCHED_NOTES),
            })
            evidence_ids = [fid]
            supported = "销售出库数量与计划一致。"

        elif cat == "split_matched":
            a = random.randint(2, planned - 2)
            b = planned - a
            t1 = pick_time(order_time, agreed_time)
            t2 = pick_time(t1, agreed_time)
            fid1 = next_flow_id(flow_seq)
            fid2 = next_flow_id(flow_seq)
            line_flows.append({
                "flow_id": fid1,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "销售出库",
                "quantity": a,
                "unit": "件",
                "time": t1,
                "status": "已过账",
                "related_flow": None,
                "source": "测试WMS",
                "note": random.choice(FIRST_BATCH_NOTES).format(a=a),
            })
            line_flows.append({
                "flow_id": fid2,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "销售出库",
                "quantity": b,
                "unit": "件",
                "time": t2,
                "status": "已过账",
                "related_flow": None,
                "source": "测试WMS",
                "note": random.choice(SECOND_BATCH_NOTES).format(b=b),
            })
            eid = next_evidence_id(evidence_seq)
            line_explanations.append({
                "evidence_id": eid,
                "order_id": order_id,
                "line_id": line_id,
                "related_flow": None,
                "time": pick_time(t2, DEADLINE),
                "record_type": "交接记录",
                "recorder": "仓库员甲",
                "content": f"本订单明细分两车装运。{fid1}发出{a}件，{fid2}发出{b}件，合计{planned}件。",
            })
            evidence_ids = [fid1, fid2, eid]
            supported = "分两批出库，合计数量与计划一致。"

        elif cat in ("quantity_difference", "insufficient_evidence"):
            diff = random.randint(1, max(1, planned // 2))
            actual = planned - diff
            t = pick_time(order_time, agreed_time)
            fid = next_flow_id(flow_seq)
            if cat == "quantity_difference":
                reason_type = random.choice(["damaged", "label", "misplaced"])
                if reason_type == "damaged":
                    ev_text = f"拣货时发现{diff}件外包装破损，移至待检区。本次先发{actual}件，剩余{diff}件待处理。"
                    note = f"本次发出{actual}件，外包装破损商品已移至待检区，见交接记录。"
                elif reason_type == "label":
                    ev_text = f"复核发现{diff}件标签贴错，已退回重新贴标。本次先发{actual}件。"
                    note = f"本次发出{actual}件，标签贴错商品已退回重贴，见复核记录。"
                else:
                    ev_text = f"装车时{diff}件被误放至相邻月台，追回后重新清点。本次实发{actual}件。"
                    note = f"本次发出{actual}件，误放月台商品已追回并重新清点，见交接记录。"
            else:
                note = random.choice(INSUFFICIENT_NOTES)
            line_flows.append({
                "flow_id": fid,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "销售出库",
                "quantity": actual,
                "unit": "件",
                "time": t,
                "status": "已过账",
                "related_flow": None,
                "source": "测试WMS",
                "note": note,
            })
            if cat == "quantity_difference":
                eid = next_evidence_id(evidence_seq)
                line_explanations.append({
                    "evidence_id": eid,
                    "order_id": order_id,
                    "line_id": line_id,
                    "related_flow": fid,
                    "time": pick_time(t, DEADLINE),
                    "record_type": "交接记录",
                    "recorder": "仓库员乙",
                    "content": ev_text,
                })
                evidence_ids = [fid, eid]
                supported = ev_text
            else:
                evidence_ids = [fid]
                supported = "现有记录不足以确定原因。"

        elif cat == "sku_mismatch":
            wrong_sku, wrong_name = random.choice([s for s in SKUS if s[0] != sku])
            t = pick_time(order_time, agreed_time)
            fid = next_flow_id(flow_seq)
            line_flows.append({
                "flow_id": fid,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": wrong_sku,
                "sku_name": wrong_name,
                "biz_type": "销售出库",
                "quantity": planned,
                "unit": "件",
                "time": t,
                "status": "已过账",
                "related_flow": None,
                "source": "测试WMS",
                "note": f"扫描出库商品为{wrong_sku}，按扫描结果登记。",
            })
            evidence_ids = [fid]
            supported = f"流水{fid}记录了{planned}件{wrong_name}（{wrong_sku}），与计划SKU不一致。"

        elif cat == "draft_only":
            t = pick_time(order_time, agreed_time)
            fid = next_flow_id(flow_seq)
            line_flows.append({
                "flow_id": fid,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "销售出库",
                "quantity": planned,
                "unit": "件",
                "time": t,
                "status": "草稿",
                "related_flow": None,
                "source": "测试WMS",
                "note": "草稿出库单已登记，待复核过账。",
            })
            eid = next_evidence_id(evidence_seq)
            line_explanations.append({
                "evidence_id": eid,
                "order_id": order_id,
                "line_id": line_id,
                "related_flow": fid,
                "time": pick_time(t, DEADLINE),
                "record_type": "复核记录",
                "recorder": "复核员甲",
                "content": f"{fid}仍为草稿状态，尚未过账，需等待仓库确认实际出库。",
            })
            evidence_ids = [fid, eid]
            supported = "存在草稿出库记录，但无已过账有效记录。"

        elif cat == "reversal_matched":
            over = random.randint(1, max(1, planned // 2))
            original_qty = planned + over
            t_orig = pick_time(order_time, agreed_time)
            t_rev = pick_time(t_orig, agreed_time)
            fid_orig = next_flow_id(flow_seq)
            fid_rev = next_flow_id(flow_seq)
            line_flows.append({
                "flow_id": fid_orig,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "销售出库",
                "quantity": original_qty,
                "unit": "件",
                "time": t_orig,
                "status": "已过账",
                "related_flow": None,
                "source": "测试WMS",
                "note": f"初录{original_qty}件，后发现多录{over}件。",
            })
            line_flows.append({
                "flow_id": fid_rev,
                "wh_order_id": f"CK-{flow_seq[0]:04d}",
                "wh_line_id": "001",
                "order_id": order_id,
                "line_id": line_id,
                "warehouse": "WH-01",
                "sku": sku,
                "sku_name": sku_name,
                "biz_type": "出库冲销",
                "quantity": over,
                "unit": "件",
                "time": t_rev,
                "status": "已过账",
                "related_flow": fid_orig,
                "source": "测试WMS",
                "note": f"冲销原流水{fid_orig}多录的{over}件。",
            })
            eid = next_evidence_id(evidence_seq)
            line_explanations.append({
                "evidence_id": eid,
                "order_id": order_id,
                "line_id": line_id,
                "related_flow": fid_rev,
                "time": pick_time(t_rev, DEADLINE),
                "record_type": "复核记录",
                "recorder": "复核员乙",
                "content": f"复核签收数量为{planned}件，原流水{fid_orig}误录为{original_qty}件。已通过{fid_rev}冲销多录的{over}件，原流水保留。",
            })
            evidence_ids = [fid_orig, fid_rev, eid]
            supported = f"原流水{fid_orig}多录{over}件，已用{fid_rev}冲销，净出库量与计划一致。"

        elif cat == "not_due":
            evidence_ids = []
            supported = f"约定出库时间为{line['agreed_time'].strftime('%Y-%m-%d %H:%M')}，尚未到截止时间，不判逾期。"

        net_shipped = planned
        diff_qty = 0
        if cat in ("quantity_difference", "insufficient_evidence"):
            net_shipped = planned - diff
            diff_qty = diff
        elif cat == "sku_mismatch":
            net_shipped = 0
            diff_qty = planned
        elif cat in ("draft_only", "not_due"):
            net_shipped = 0
            diff_qty = planned

        case = {
            "order_id": order_id,
            "order_line_id": line_id,
            "question": "这条订单明细到截止时间还差多少，现有资料如何解释？",
            "expected_quantity": planned,
            "matched_net_shipped_quantity": net_shipped,
            "difference_quantity": diff_qty,
            "expected_finding": cat,
            "evidence_ids": evidence_ids,
            "supported_explanation": supported,
            "must_not_claim": MUST_NOT_CLAIM.get(cat, []),
        }
        cases.append(case)
        flows.extend(line_flows)
        explanations.extend(line_explanations)

    # Add a few normal-record noise explanations to keep RAG from overfitting.
    case_map = {(c["order_id"], c["order_line_id"]): c for c in cases}
    normal_lines = [l for l in lines if l["category"] == "matched"]
    random.shuffle(normal_lines)
    for line in normal_lines[:8]:
        t = pick_time(line["order_time"], DEADLINE)
        eid = next_evidence_id(evidence_seq)
        explanations.append({
            "evidence_id": eid,
            "order_id": line["order_id"],
            "line_id": line["line_id"],
            "related_flow": None,
            "time": t,
            "record_type": "交接记录",
            "recorder": "仓库员丙",
            "content": f"{line['order_id']}/{line['line_id']}按拣货单完成交接，单据与实物核对一致。",
        })
        case_map[(line["order_id"], line["line_id"])]["evidence_ids"].append(eid)

    return flows, explanations, cases


def write_excel(lines, flows, explanations):
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_plan = wb.create_sheet("发货计划")
    ws_plan.append(PLAN_HEADERS)
    for line in lines:
        ws_plan.append([
            line["order_id"],
            line["line_id"],
            "WH-01",
            line["sku"],
            line["sku_name"],
            line["planned"],
            "件",
            line["order_time"],
            line["agreed_time"],
            "生效",
            "测试OMS",
            line["plan_note"] or None,
        ])

    ws_flow = wb.create_sheet("仓储流水")
    ws_flow.append(FLOW_HEADERS)
    for f in flows:
        ws_flow.append([
            f["flow_id"],
            f["wh_order_id"],
            f["wh_line_id"],
            f["order_id"],
            f["line_id"],
            f["warehouse"],
            f["sku"],
            f["sku_name"],
            f["biz_type"],
            f["quantity"],
            f["unit"],
            f["time"],
            f["status"],
            f["related_flow"] or None,
            f["source"],
            f["note"] or None,
        ])

    ws_exp = wb.create_sheet("业务说明")
    ws_exp.append(EXPLANATION_HEADERS)
    for e in explanations:
        ws_exp.append([
            e["evidence_id"],
            e["order_id"],
            e["line_id"],
            e["related_flow"] or None,
            e["time"],
            e["record_type"],
            e["recorder"],
            e["content"],
        ])

    wb.save(OUTPUT_XLSX)


def write_json(cases):
    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    cases_sorted = sorted(cases, key=lambda c: (c["order_id"], c["order_line_id"]))
    payload = {
        "as_of": "2026-09-09T20:00:00+08:00",
        "timezone": "Asia/Shanghai",
        "cases": cases_sorted,
    }
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def verify():
    errors = []
    wb = openpyxl.load_workbook(OUTPUT_XLSX, data_only=True)

    ws_plan = wb["发货计划"]
    ws_flow = wb["仓储流水"]
    ws_exp = wb["业务说明"]

    plan_headers = [ws_plan.cell(1, c).value for c in range(1, ws_plan.max_column + 1)]
    flow_headers = [ws_flow.cell(1, c).value for c in range(1, ws_flow.max_column + 1)]
    exp_headers = [ws_exp.cell(1, c).value for c in range(1, ws_exp.max_column + 1)]

    if plan_headers != PLAN_HEADERS:
        errors.append(f"发货计划表头不一致: {plan_headers}")
    if flow_headers != FLOW_HEADERS:
        errors.append(f"仓储流水表头不一致: {flow_headers}")
    if exp_headers != EXPLANATION_HEADERS:
        errors.append(f"业务说明表头不一致: {exp_headers}")

    plans = []
    plan_keys = set()
    for r in range(2, ws_plan.max_row + 1):
        p = {
            "order_id": str(ws_plan.cell(r, 1).value),
            "line_id": str(ws_plan.cell(r, 2).value),
            "warehouse": ws_plan.cell(r, 3).value,
            "sku": ws_plan.cell(r, 4).value,
            "sku_name": ws_plan.cell(r, 5).value,
            "planned": ws_plan.cell(r, 6).value,
            "unit": ws_plan.cell(r, 7).value,
            "order_time": ws_plan.cell(r, 8).value,
            "agreed_time": ws_plan.cell(r, 9).value,
            "order_status": ws_plan.cell(r, 10).value,
            "source": ws_plan.cell(r, 11).value,
            "note": ws_plan.cell(r, 12).value,
        }
        plans.append(p)
        key = (p["order_id"], p["line_id"])
        if key in plan_keys:
            errors.append(f"重复的发货计划键: {key}")
        plan_keys.add(key)

    flows = []
    flow_ids = set()
    flow_by_id = {}
    for r in range(2, ws_flow.max_row + 1):
        f = {
            "flow_id": ws_flow.cell(r, 1).value,
            "wh_order_id": ws_flow.cell(r, 2).value,
            "wh_line_id": ws_flow.cell(r, 3).value,
            "order_id": ws_flow.cell(r, 4).value,
            "line_id": ws_flow.cell(r, 5).value,
            "warehouse": ws_flow.cell(r, 6).value,
            "sku": ws_flow.cell(r, 7).value,
            "sku_name": ws_flow.cell(r, 8).value,
            "biz_type": ws_flow.cell(r, 9).value,
            "quantity": ws_flow.cell(r, 10).value,
            "unit": ws_flow.cell(r, 11).value,
            "time": ws_flow.cell(r, 12).value,
            "status": ws_flow.cell(r, 13).value,
            "related_flow": ws_flow.cell(r, 14).value,
            "source": ws_flow.cell(r, 15).value,
            "note": ws_flow.cell(r, 16).value,
        }
        flows.append(f)
        if f["flow_id"] in flow_ids:
            errors.append(f"重复的流水号: {f['flow_id']}")
        flow_ids.add(f["flow_id"])
        flow_by_id[f["flow_id"]] = f

    explanations = []
    evidence_ids = set()
    for r in range(2, ws_exp.max_row + 1):
        e = {
            "evidence_id": ws_exp.cell(r, 1).value,
            "order_id": ws_exp.cell(r, 2).value,
            "line_id": ws_exp.cell(r, 3).value,
            "related_flow": ws_exp.cell(r, 4).value,
            "time": ws_exp.cell(r, 5).value,
            "record_type": ws_exp.cell(r, 6).value,
            "recorder": ws_exp.cell(r, 7).value,
            "content": ws_exp.cell(r, 8).value,
        }
        explanations.append(e)
        if e["evidence_id"] in evidence_ids:
            errors.append(f"重复的证据编号: {e['evidence_id']}")
        evidence_ids.add(e["evidence_id"])

    # Row counts
    if len(plans) != 150:
        errors.append(f"发货计划行数错误: {len(plans)}")
    if len(flows) != 170:
        errors.append(f"仓储流水行数错误: {len(flows)}")
    if not (35 <= len(explanations) <= 50):
        errors.append(f"业务说明行数错误: {len(explanations)}")

    # Foreign keys
    for f in flows:
        if (f["order_id"], f["line_id"]) not in plan_keys:
            errors.append(f"流水外键不存在: {f['flow_id']} -> {(f['order_id'], f['line_id'])}")
        if f["related_flow"] and f["related_flow"] not in flow_by_id:
            errors.append(f"关联流水号不存在: {f['flow_id']} -> {f['related_flow']}")
    for e in explanations:
        if (e["order_id"], e["line_id"]) not in plan_keys:
            errors.append(f"说明外键不存在: {e['evidence_id']} -> {(e['order_id'], e['line_id'])}")
        if e["related_flow"] and e["related_flow"] not in flow_by_id:
            errors.append(f"说明关联流水不存在: {e['evidence_id']} -> {e['related_flow']}")

    # Time constraints
    for p in plans:
        if p["order_time"] < ORDER_START or p["order_time"] > ORDER_END:
            errors.append(f"下单时间越界: {p['order_id']}/{p['line_id']}")
        if p["agreed_time"] < p["order_time"]:
            errors.append(f"约定出库时间早于下单时间: {p['order_id']}/{p['line_id']}")
        if p["agreed_time"] > DEADLINE and p["agreed_time"] < NOT_DUE_START:
            errors.append(f"约定出库时间在截止后但早于not_due区间: {p['order_id']}/{p['line_id']}")
    for f in flows:
        if f["time"] > DEADLINE:
            errors.append(f"流水时间晚于截止时间: {f['flow_id']}")
        if f["time"] < datetime.datetime(2026, 9, 1, 0, 0, 0):
            errors.append(f"流水时间过早: {f['flow_id']}")
    for e in explanations:
        if e["time"] > DEADLINE:
            errors.append(f"说明时间晚于截止时间: {e['evidence_id']}")

    # Reversal rules
    reversal_sums = collections.defaultdict(int)
    for f in flows:
        if f["biz_type"] == "出库冲销":
            orig = flow_by_id.get(f["related_flow"])
            if not orig:
                errors.append(f"冲销无原流水: {f['flow_id']}")
                continue
            if orig["status"] != "已过账":
                errors.append(f"冲销指向非已过账流水: {f['flow_id']} -> {orig['flow_id']}")
            if orig["biz_type"] != "销售出库":
                errors.append(f"冲销指向非销售出库: {f['flow_id']} -> {orig['flow_id']}")
            if not ((orig["order_id"], orig["line_id"], orig["warehouse"], orig["sku"], orig["unit"]) ==
                    (f["order_id"], f["line_id"], f["warehouse"], f["sku"], f["unit"])):
                errors.append(f"冲销事务字段不一致: {f['flow_id']} -> {orig['flow_id']}")
            if f["time"] < orig["time"]:
                errors.append(f"冲销时间早于原事件: {f['flow_id']} -> {orig['flow_id']}")
            reversal_sums[orig["flow_id"]] += f["quantity"]
    for orig_id, total in reversal_sums.items():
        orig = flow_by_id[orig_id]
        if total > orig["quantity"]:
            errors.append(f"冲销累计超原数量: {orig_id} 原={orig['quantity']} 冲={total}")

    # Per-plan recomputation vs JSON
    with open(OUTPUT_JSON, "r", encoding="utf-8") as f:
        answer = json.load(f)
    cases = answer["cases"]
    if len(cases) != 150:
        errors.append(f"JSON cases 数量错误: {len(cases)}")

    flows_by_plan = collections.defaultdict(list)
    for f in flows:
        flows_by_plan[(f["order_id"], f["line_id"])].append(f)
    exps_by_plan = collections.defaultdict(list)
    for e in explanations:
        exps_by_plan[(e["order_id"], e["line_id"])].append(e)

    case_map = {(c["order_id"], c["order_line_id"]): c for c in cases}

    for p in plans:
        key = (p["order_id"], p["line_id"])
        if key not in case_map:
            errors.append(f"JSON 缺少 case: {key}")
            continue
        case = case_map[key]
        plan_flows = flows_by_plan.get(key, [])
        plan_exps = exps_by_plan.get(key, [])

        posted = [f for f in plan_flows if f["status"] == "已过账" and f["time"] <= DEADLINE]
        drafts = [f for f in plan_flows if f["status"] == "草稿"]

        if p["agreed_time"] > DEADLINE:
            expected_finding = "not_due"
            net = 0
            diff = p["planned"]
        else:
            mismatch = [f for f in posted if (f["sku"], f["warehouse"], f["unit"]) != (p["sku"], p["warehouse"], p["unit"])]
            if mismatch:
                expected_finding = "sku_mismatch"
                net = 0
                diff = p["planned"]
            else:
                net = sum(f["quantity"] for f in posted if f["biz_type"] == "销售出库") - \
                      sum(f["quantity"] for f in posted if f["biz_type"] == "出库冲销")
                diff = p["planned"] - net
                if diff == 0:
                    revs = [f for f in posted if f["biz_type"] == "出库冲销"]
                    sales = [f for f in posted if f["biz_type"] == "销售出库"]
                    if revs:
                        expected_finding = "reversal_matched"
                    elif len(sales) > 1:
                        expected_finding = "split_matched"
                    else:
                        expected_finding = "matched"
                else:
                    if drafts:
                        expected_finding = "draft_only"
                    elif plan_exps:
                        expected_finding = "quantity_difference"
                    else:
                        expected_finding = "insufficient_evidence"

        if case["expected_quantity"] != p["planned"]:
            errors.append(f"{key} expected_quantity 不一致: json={case['expected_quantity']} plan={p['planned']}")
        if case["matched_net_shipped_quantity"] != net:
            errors.append(f"{key} matched_net_shipped_quantity 不一致: json={case['matched_net_shipped_quantity']} calc={net}")
        if case["difference_quantity"] != diff:
            errors.append(f"{key} difference_quantity 不一致: json={case['difference_quantity']} calc={diff}")
        if case["expected_finding"] != expected_finding:
            errors.append(f"{key} expected_finding 不一致: json={case['expected_finding']} calc={expected_finding}")

        # Evidence ids consistency
        relevant_flow_ids = [f["flow_id"] for f in plan_flows if f["status"] in ("已过账", "草稿") and f["time"] <= DEADLINE]
        expected_evidence = sorted(set(relevant_flow_ids + [e["evidence_id"] for e in plan_exps]))
        if sorted(case["evidence_ids"]) != expected_evidence:
            errors.append(f"{key} evidence_ids 不一致: json={sorted(case['evidence_ids'])} calc={expected_evidence}")

    # Scenario distribution
    dist = collections.Counter(c["expected_finding"] for c in cases)
    for cat, count in REQUIRED_DISTRIBUTION.items():
        if dist.get(cat, 0) != count:
            errors.append(f"场景分布不符: {cat} 应为 {count}，实际 {dist.get(cat, 0)}")

    return errors, plans, flows, explanations, cases


def print_spot_checks(plans, flows, explanations, cases):
    targets = ["split_matched", "reversal_matched", "sku_mismatch", "insufficient_evidence"]
    case_map = {(c["order_id"], c["order_line_id"]): c for c in cases}
    flows_by_plan = collections.defaultdict(list)
    for f in flows:
        flows_by_plan[(f["order_id"], f["line_id"])].append(f)
    exps_by_plan = collections.defaultdict(list)
    for e in explanations:
        exps_by_plan[(e["order_id"], e["line_id"])].append(e)
    plan_map = {(p["order_id"], p["line_id"]): p for p in plans}

    print("\n=== 独立抽查（按类别各一条，从输出文件重新计算）===")
    for cat in targets:
        for c in cases:
            if c["expected_finding"] == cat:
                key = (c["order_id"], c["order_line_id"])
                p = plan_map[key]
                print(f"\n类别: {cat}")
                print(f"  订单/明细: {key[0]}/{key[1]}")
                print(f"  计划: SKU={p['sku']}, 应发={p['planned']}")
                for f in sorted(flows_by_plan[key], key=lambda x: x["time"]):
                    rel = f" -> {f['related_flow']}" if f["related_flow"] else ""
                    print(f"  流水: {f['flow_id']} {f['biz_type']} {f['status']} SKU={f['sku']} qty={f['quantity']} time={f['time']}{rel}")
                for e in exps_by_plan[key]:
                    print(f"  说明: {e['evidence_id']} time={e['time']} content={e['content'][:60]}...")
                print(f"  答案JSON: net={c['matched_net_shipped_quantity']} diff={c['difference_quantity']} finding={c['expected_finding']}")
                print(f"  核对: evidence_ids={c['evidence_ids']}")
                break


def main():
    random.seed(20240909)

    if "--verify-only" in sys.argv:
        errors, plans, flows, explanations, cases = verify()
        if errors:
            print("验证失败:")
            for e in errors:
                print("  -", e)
            sys.exit(1)
        print("验证全部通过（仅验证模式）")
        print_spot_checks(plans, flows, explanations, cases)
        return

    lines = build_lines()
    flows, explanations, cases = generate_flows_and_answers(lines)
    write_excel(lines, flows, explanations)
    write_json(cases)
    print(f"已生成: {OUTPUT_XLSX}")
    print(f"已生成: {OUTPUT_JSON}")
    print(f"发货计划: {len(lines)} 行, 仓储流水: {len(flows)} 行, 业务说明: {len(explanations)} 行")

    errors, plans, flows, explanations, cases = verify()
    if errors:
        print("\n验证失败:")
        for e in errors:
            print("  -", e)
        sys.exit(1)

    print("\n=== 程序化验证全部通过 ===")
    print_spot_checks(plans, flows, explanations, cases)


if __name__ == "__main__":
    main()
