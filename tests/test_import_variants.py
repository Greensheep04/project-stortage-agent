import copy
from pathlib import Path

import openpyxl
import pytest

from pi_market.db import drop_dataset
from pi_market.importer import import_excel

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REAL_XLSX = PROJECT_ROOT / "deliverables" / "仓储异常测试数据.xlsx"
AS_OF = "2026-09-09 20:00:00"


def _load_workbook():
    return openpyxl.load_workbook(REAL_XLSX, data_only=False)


def _copy_and_import(variant_path: Path, dataset_id: str) -> dict:
    drop_dataset(dataset_id)
    return import_excel(variant_path, AS_OF)


def _swap_column_values(ws, col_a: int, col_b: int):
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=col_a).value
        b = ws.cell(row=r, column=col_b).value
        ws.cell(row=r, column=col_a).value = b
        ws.cell(row=r, column=col_b).value = a


def _rotate_columns(ws, cols: list):
    """Rotate values among the given 1-based columns."""
    for r in range(1, ws.max_row + 1):
        values = [ws.cell(row=r, column=c).value for c in cols]
        rotated = values[-1:] + values[:-1]
        for c, v in zip(cols, rotated):
            ws.cell(row=r, column=c).value = v


def _make_reordered(tmp_path: Path) -> Path:
    wb = _load_workbook()
    # 发货计划：swap 订单号 (1) and SKU (4).
    _swap_column_values(wb["发货计划"], 1, 4)
    # 仓储流水：rotate first three columns 1->2->3->1.
    _rotate_columns(wb["仓储流水"], [1, 2, 3])
    path = tmp_path / "reordered.xlsx"
    wb.save(path)
    return path


def _make_aliased(tmp_path: Path) -> Path:
    wb = _load_workbook()
    plan = wb["发货计划"]
    plan.cell(row=1, column=1, value="订单编号")
    plan.cell(row=1, column=6, value="应发件数")
    plan.cell(row=1, column=8, value="下单日期")

    flow = wb["仓储流水"]
    flow.cell(row=1, column=1, value="流水ID")
    flow.cell(row=1, column=9, value="业务类型")
    flow.cell(row=1, column=13, value="记录状态")

    note = wb["业务说明"]
    note.cell(row=1, column=1, value="证据ID")
    note.cell(row=1, column=8, value="说明")

    path = tmp_path / "aliased.xlsx"
    wb.save(path)
    return path


def _make_title_row(tmp_path: Path) -> Path:
    wb = _load_workbook()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value=f"{sheet_name}数据")
    path = tmp_path / "title_row.xlsx"
    wb.save(path)
    return path


def _make_missing_required(tmp_path: Path) -> Path:
    wb = _load_workbook()
    plan = wb["发货计划"]
    plan["A2"] = None  # order_id missing; ws.cell(..., value=None) does not clear
    path = tmp_path / "missing_required.xlsx"
    wb.save(path)
    return path


def _make_duplicate_plan_key(tmp_path: Path) -> Path:
    wb = _load_workbook()
    plan = wb["发货计划"]
    row2 = [plan.cell(row=2, column=c).value for c in range(1, plan.max_column + 1)]
    for c, v in enumerate(row2, start=1):
        plan.cell(row=3, column=c, value=v)
    path = tmp_path / "duplicate_plan_key.xlsx"
    wb.save(path)
    return path


def _make_reversal_ref_draft(tmp_path: Path) -> Path:
    wb = _load_workbook()
    flow = wb["仓储流水"]
    # Find a reversal row and set its referenced original row to draft.
    for r in range(2, flow.max_row + 1):
        if flow.cell(row=r, column=9).value == "出库冲销":
            ref = flow.cell(row=r, column=14).value
            break
    else:
        raise RuntimeError("no reversal row found")
    for r2 in range(2, flow.max_row + 1):
        if flow.cell(row=r2, column=1).value == ref:
            flow.cell(row=r2, column=13, value="草稿")
            break
    path = tmp_path / "reversal_ref_draft.xlsx"
    wb.save(path)
    return path


def _make_reversal_time_before(tmp_path: Path) -> Path:
    wb = _load_workbook()
    flow = wb["仓储流水"]
    for r in range(2, flow.max_row + 1):
        if flow.cell(row=r, column=9).value == "出库冲销":
            flow.cell(row=r, column=12, value=flow.cell(row=r, column=12).value.replace(year=2024))
            break
    path = tmp_path / "reversal_time_before.xlsx"
    wb.save(path)
    return path


def _make_reversal_over(tmp_path: Path) -> Path:
    wb = _load_workbook()
    flow = wb["仓储流水"]
    for r in range(2, flow.max_row + 1):
        if flow.cell(row=r, column=9).value == "出库冲销":
            flow.cell(row=r, column=10, value=9999)
            break
    path = tmp_path / "reversal_over.xlsx"
    wb.save(path)
    return path


def test_reordered_columns(tmp_path):
    path = _make_reordered(tmp_path)
    result = _copy_and_import(path, "ds-reordered")
    assert result["status"] == "ready"
    assert result["plan_rows"] == 150
    assert result["flow_rows"] == 170
    assert result["note_rows"] == 48


def test_aliased_headers(tmp_path):
    path = _make_aliased(tmp_path)
    result = _copy_and_import(path, "ds-aliased")
    assert result["status"] == "ready"
    assert result["plan_rows"] == 150


def test_top_title_row(tmp_path):
    path = _make_title_row(tmp_path)
    result = _copy_and_import(path, "ds-title")
    assert result["status"] == "ready"
    assert result["plan_rows"] == 150


def test_missing_required(tmp_path):
    path = _make_missing_required(tmp_path)
    result = _copy_and_import(path, "ds-missing")
    assert result["status"] == "failed"
    assert any(i["field"] == "order_id" and "必填" in i["message"] for i in result["issues"])


def test_duplicate_plan_key(tmp_path):
    path = _make_duplicate_plan_key(tmp_path)
    result = _copy_and_import(path, "ds-dup-plan")
    assert result["status"] == "failed"
    assert any("重复" in i["message"] for i in result["issues"])


def test_reversal_ref_draft(tmp_path):
    path = _make_reversal_ref_draft(tmp_path)
    result = _copy_and_import(path, "ds-rev-draft")
    assert result["status"] == "failed"
    assert any("冲销" in i["message"] and "已过账" in i["message"] for i in result["issues"])


def test_reversal_time_before(tmp_path):
    path = _make_reversal_time_before(tmp_path)
    result = _copy_and_import(path, "ds-rev-time")
    assert result["status"] == "failed"
    assert any("冲销时间早于" in i["message"] for i in result["issues"])


def test_reversal_over(tmp_path):
    path = _make_reversal_over(tmp_path)
    result = _copy_and_import(path, "ds-rev-over")
    assert result["status"] == "failed"
    assert any("冲销累计超原数量" in i["message"] for i in result["issues"])
